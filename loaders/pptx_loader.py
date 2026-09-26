# loaders/pptx_loader.py
from pathlib import Path
from typing import List
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from captioning import AbstractCaptioner
from vector_db.chunk import Chunk
from . import AbstractLoader
from .text_utils import split_into_chunks

EXTRACTED_IMAGE_DIR = Path("data/extracted_images")

class PPTXLoader(AbstractLoader):
    def __init__(self, captioner: AbstractCaptioner):
        self.captioner = captioner

    def load(self, path: Path, chunk_size: int, max_chars: int | None = None) -> List[Chunk]:
        prs = Presentation(str(path))
        chunks = []
        for slide_num, slide in enumerate(prs.slides):
            texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
            slide_text = "\n".join(texts)
            if not slide_text.strip():
                continue
            for i, piece in enumerate(split_into_chunks(text=slide_text, max_sentences=chunk_size, max_chars=max_chars)):
                chunks.append(Chunk(
                    doc_id=path.stem, 
                    chunk_id=f"s{slide_num}_c{i}",
                    text=piece
                ))
            chunks.extend(self._extract_images(slide, slide_num, path, slide_text))
        return chunks

    def _extract_images(self, slide, slide_num: int, source_path: Path, slide_text: str) -> List[Chunk]:
        chunks = []
        EXTRACTED_IMAGE_DIR.mkdir(parents=True, exist_ok=True)

        for shape_index, shape in enumerate(slide.shapes):
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            image = shape.image
            image_path = EXTRACTED_IMAGE_DIR / f"{source_path.stem}_s{slide_num}_i{shape_index}.{image.ext}"
            image_path.write_bytes(image.blob)

            caption = self.captioner.caption(image_path, context=slide_text.strip() or None)
            chunks.append(Chunk(
                doc_id=source_path.stem, chunk_id=f"s{slide_num}_img{shape_index}",
                text=caption, modality="image",
                source_path=source_path, page_num=slide_num, image_path=image_path,
            ))
        return chunks